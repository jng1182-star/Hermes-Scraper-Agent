"""
Hermes — HTTP server (stdlib only, no FastAPI/uvicorn).
Serves the dashboard at /dashboard/ and provides a JSON API.
Runs the HTTP server in a daemon thread; main thread stays free for signals.
"""
import cgi
import json
import os
import threading
import time
from collections import deque
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from socketserver import ThreadingMixIn
from typing import Optional

# ── Lock all LLM routing to local Ollama before any crewai/litellm import ────
# CrewAI's llm_utils.py falls back to OPENAI_API_BASE / OPENAI_BASE_URL / BASE_URL
# if those env vars exist — stale ngrok/cloud values from other tools override our
# explicit base_url and cause ERR_NGROK_3004 errors. Nuke them here at process start.
_ollama_host = os.getenv("OLLAMA_HOST", "http://127.0.0.1:11434").rstrip("/")
_ollama_v1   = _ollama_host + "/v1"
os.environ["OLLAMA_HOST"]             = _ollama_host
os.environ["OPENAI_API_BASE"]         = _ollama_v1
os.environ["OPENAI_BASE_URL"]         = _ollama_v1
os.environ["CREWAI_TELEMETRY_OPT_OUT"] = "true"
for _v in ("BASE_URL", "API_BASE", "HERMES_TUNNEL_TOKEN"):
    os.environ.pop(_v, None)

HOST = os.getenv("HOST", "0.0.0.0")
PORT = int(os.getenv("PORT", "8000"))

# ── Shared run state ──────────────────────────────────────────────────────────
_run_state = {
    "running":      False,
    "logs":         deque(maxlen=300),
    "error":        None,
    "agent_states": {},
    "timed_out":    False,
    "retry_count":  0,
    "start_ts":     None,
}
_state_lock = threading.Lock()
_stop_flag  = threading.Event()   # set to request graceful abort
_worker_thread_id: Optional[int] = None
_worker_thread: Optional[threading.Thread] = None   # reference for join/kill


def _kill_ollama_now():
    """Kill the Ollama runner process to unblock any in-flight LLM call."""
    import subprocess
    try:
        r = subprocess.run(["pgrep", "-f", "ollama runner"], capture_output=True, text=True)
        for pid_str in r.stdout.strip().splitlines():
            try:
                os.kill(int(pid_str), 9)   # SIGKILL — instant
            except (ValueError, ProcessLookupError):
                pass
    except Exception:
        pass
    # Also stop model via ollama CLI (best-effort)
    try:
        subprocess.run(["ollama", "stop", "gemma4:e4b"],  capture_output=True, timeout=3)
        subprocess.run(["ollama", "stop", "gemma4:26b"], capture_output=True, timeout=3)
    except Exception:
        pass

# Map CrewAI Agent.role → dashboard node IDs
_ROLE_TO_NODE = {
    "social data scraper": "scraper",
    "engagement analyst":  "analyst",
    "intelligence reporter": "reporter",
}


def _set_agent_active(node_id: str, label: str = ""):
    with _state_lock:
        for aid, s in _run_state["agent_states"].items():
            if s == "active":
                _run_state["agent_states"][aid] = "done"
        _run_state["agent_states"][node_id] = "active"
        _run_state["logs"].append(f"[Agent] {label or node_id} → active")


def _set_agent_done(node_id: str, label: str = ""):
    with _state_lock:
        _run_state["agent_states"][node_id] = "done"
        _run_state["logs"].append(f"[Agent] {label or node_id} → done")


def _patch_crewai_agent():
    """Monkey-patch Agent.execute_task to fire state hooks synchronously."""
    from crewai.agent.core import Agent as _Agent
    _orig = _Agent.execute_task

    def _patched(self, task, context=None, tools=None):
        role = (getattr(self, "role", "") or "").lower().strip()
        node_id = _ROLE_TO_NODE.get(role)
        if node_id:
            _set_agent_active(node_id, self.role)
        result = _orig(self, task, context=context, tools=tools)
        if node_id:
            _set_agent_done(node_id, self.role)
        return result

    _Agent.execute_task = _patched


_patch_crewai_agent()


# ── Partial results helper ────────────────────────────────────────────────────

def _try_partial_report() -> Optional[dict]:
    """Return whatever structured competitor data is available from checkpoints.
    Priority: reporter cp → analyst cp → scraper cp (raw, best-effort parse).
    Returns None if nothing parseable is available yet."""
    cp_dir = Path("data/checkpoints")
    for phase in ("reporter", "analyst", "scraper"):
        cp = cp_dir / f"{phase}.json"
        if not cp.exists():
            continue
        try:
            raw = json.loads(cp.read_text(encoding="utf-8")).get("output", "")
            # Try to extract a competitors array from the text
            import re as _re
            # Strip markdown code fences
            cleaned = _re.sub(r"```(?:json)?\s*", "", raw)
            cleaned = _re.sub(r"```", "", cleaned).strip()
            start = cleaned.find("{")
            end   = cleaned.rfind("}") + 1
            if start != -1 and end > start:
                data = json.loads(cleaned[start:end])
                if "competitors" in data and data["competitors"]:
                    data["_partial_phase"] = phase
                    return data
            # Try array
            start = cleaned.find("[")
            end   = cleaned.rfind("]") + 1
            if start != -1 and end > start:
                lst = json.loads(cleaned[start:end])
                if lst:
                    return {"competitors": lst, "_partial_phase": phase}
        except Exception:
            continue
    return None


# ── Pipeline runner ───────────────────────────────────────────────────────────

def _pipeline_hook(node_id: str, state: str):
    """Called from main.py for gate state and stall events."""
    if node_id == "__stall__":
        retry_n = int(state.split(":")[1]) if ":" in state else 1
        with _state_lock:
            _run_state["timed_out"]   = True
            _run_state["retry_count"] = retry_n
            for aid in _run_state["agent_states"]:
                _run_state["agent_states"][aid] = "idle"
            _run_state["logs"].append(
                f"[WATCHDOG] Stall detected — retry {retry_n}. Restarting…"
            )
    else:
        with _state_lock:
            _run_state["agent_states"][node_id] = state
            _run_state["logs"].append(
                f"[Gate] Approval Gate → {state}"
            )


def _run_worker(params: dict):
    import sys
    import io
    import ctypes
    import main as _main

    global _worker_thread_id, _worker_thread
    _stop_flag.clear()
    _worker_thread_id  = threading.current_thread().ident
    _worker_thread     = threading.current_thread()

    with _state_lock:
        _run_state.update(
            running=True, error=None, timed_out=False, retry_count=0,
            start_ts=time.monotonic(),
            agent_states={"scraper": "idle", "analyst": "idle",
                          "reporter": "idle", "gate": "idle"},
        )
        _run_state["logs"].clear()

    class _Cap(io.TextIOBase):
        def write(self, s):
            if s.strip():
                with _state_lock:
                    _run_state["logs"].append(s.rstrip())
                # Check stop flag on every write
                if _stop_flag.is_set():
                    raise RuntimeError("Analysis stopped by user.")
            return len(s)
        def flush(self): pass

    old_out, old_err = sys.stdout, sys.stderr
    sys.stdout = sys.stderr = _Cap()
    _main._state_hook = _pipeline_hook
    try:
        _main.run_pipeline(params)
        with _state_lock:
            for aid, s in _run_state["agent_states"].items():
                if s == "active":
                    _run_state["agent_states"][aid] = "done"
    except (RuntimeError, SystemExit) as e:
        if _stop_flag.is_set() or "stopped by user" in str(e).lower():
            with _state_lock:
                _run_state["logs"].append("[STOP] Analysis cancelled by user.")
                for aid in _run_state["agent_states"]:
                    _run_state["agent_states"][aid] = "idle"
        else:
            with _state_lock:
                _run_state["error"] = str(e)
                _run_state["logs"].append(f"ERROR: {e}")
    except Exception as e:
        if _stop_flag.is_set():
            with _state_lock:
                _run_state["logs"].append("[STOP] Analysis cancelled by user.")
                for aid in _run_state["agent_states"]:
                    _run_state["agent_states"][aid] = "idle"
        else:
            with _state_lock:
                _run_state["error"] = str(e)
                _run_state["logs"].append(f"ERROR: {e}")
    finally:
        sys.stdout, sys.stderr = old_out, old_err
        _main._state_hook = None
        _worker_thread_id = None
        _worker_thread    = None
        _stop_flag.clear()
        with _state_lock:
            _run_state["running"] = False


# ── HTTP server ───────────────────────────────────────────────────────────────

DASHBOARD_DIR = Path(__file__).parent / "dashboard"


def _extract_file_text(path: Path, ext: str, raw_bytes: bytes) -> str:
    """Best-effort text extraction from uploaded files. Gracefully degrades."""
    try:
        if ext in (".txt", ".csv", ".md", ".json"):
            return raw_bytes.decode("utf-8", errors="replace")

        if ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                import io
                wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
                lines = []
                for sheet in wb.worksheets:
                    lines.append(f"=== Sheet: {sheet.title} ===")
                    for row in sheet.iter_rows(values_only=True):
                        row_vals = [str(c) if c is not None else "" for c in row]
                        if any(v.strip() for v in row_vals):
                            lines.append("\t".join(row_vals))
                return "\n".join(lines)
            except ImportError:
                pass
            # Fallback: try xlrd for .xls
            try:
                import xlrd, io
                wb = xlrd.open_workbook(file_contents=raw_bytes)
                lines = []
                for sheet in wb.sheets():
                    lines.append(f"=== Sheet: {sheet.name} ===")
                    for rx in range(sheet.nrows):
                        lines.append("\t".join(str(sheet.cell_value(rx, cx)) for cx in range(sheet.ncols)))
                return "\n".join(lines)
            except ImportError:
                return f"[Excel file: {path.name} — install openpyxl to extract content]"

        if ext == ".pdf":
            try:
                import pdfplumber, io
                text_parts = []
                with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
                    for page in pdf.pages:
                        t = page.extract_text()
                        if t:
                            text_parts.append(t)
                return "\n".join(text_parts)
            except ImportError:
                pass
            try:
                import PyPDF2, io
                reader = PyPDF2.PdfReader(io.BytesIO(raw_bytes))
                return "\n".join(p.extract_text() or "" for p in reader.pages)
            except ImportError:
                return f"[PDF file: {path.name} — install pdfplumber to extract content]"

        if ext in (".docx",):
            try:
                import docx, io
                doc = docx.Document(io.BytesIO(raw_bytes))
                return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            except ImportError:
                return f"[DOCX file: {path.name} — install python-docx to extract content]"

        if ext in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
            # Store image — agents can reference it by filename
            return f"[Image file: {path.name} — {len(raw_bytes):,} bytes — visual reference only]"

    except Exception as e:
        return f"[Extraction failed for {path.name}: {e}]"

    return ""


_SAVED_RUNS_PATH = Path("data/saved_runs.json")
_saved_runs_lock = threading.Lock()

def _load_saved_runs() -> dict:
    with _saved_runs_lock:
        try:
            if _SAVED_RUNS_PATH.exists():
                return json.loads(_SAVED_RUNS_PATH.read_text(encoding="utf-8"))
        except Exception:
            pass
        return {"runs": []}

def _save_saved_runs(payload: dict):
    with _saved_runs_lock:
        _SAVED_RUNS_PATH.parent.mkdir(parents=True, exist_ok=True)
        _SAVED_RUNS_PATH.write_text(json.dumps(payload), encoding="utf-8")


class _ThreadingServer(ThreadingMixIn, HTTPServer):
    daemon_threads = True


class Handler(BaseHTTPRequestHandler):

    def do_GET(self):
        p = self.path.split("?")[0]

        if p == "/status":
            report_path = Path("data/report.json")
            with _state_lock:
                start   = _run_state["start_ts"]
                running = _run_state["running"]
                elapsed = round(time.monotonic() - start) if start else 0
                payload = {
                    "running":      running,
                    "report_ready": report_path.exists(),
                    "error":        _run_state["error"],
                    "logs":         list(_run_state["logs"])[-100:],
                    "agent_states": dict(_run_state["agent_states"]),
                    "timed_out":    _run_state["timed_out"],
                    "retry_count":  _run_state["retry_count"],
                    "elapsed_secs": elapsed,
                    "partial_report": _try_partial_report() if running else None,
                }
            self._json(200, payload)

        elif p == "/debug":
            import subprocess
            ollama_host = os.getenv("OLLAMA_HOST", "NOT SET")
            token = os.getenv("HERMES_TUNNEL_TOKEN", "").strip()
            # Test connectivity to Ollama
            try:
                import urllib.request
                req = urllib.request.Request(
                    ollama_host.rstrip("/") + "/api/version",
                    headers={"X-Hermes-Token": token} if token else {},
                )
                with urllib.request.urlopen(req, timeout=5) as r:
                    ollama_status = f"OK {r.status}"
                    ollama_body   = r.read(100).decode()
            except Exception as e:
                ollama_status = f"FAIL: {e}"
                ollama_body   = ""
            self._json(200, {
                "OLLAMA_HOST":         ollama_host,
                "HERMES_TOKEN_SET":    bool(token),
                "OLLAMA_BASE_URL":     ollama_host.rstrip("/") + "/v1",
                "ollama_connectivity": ollama_status,
                "ollama_response":     ollama_body,
            })

        elif p == "/report":
            rp = Path("data/report.json")
            if rp.exists():
                self._json(200, json.loads(rp.read_text()))
            else:
                self._json(200, {})

        elif p == "/saved-runs":
            self._json(200, _load_saved_runs())

        elif p == "/uploaded-files":
            up_dir = Path("data/uploads")
            files = []
            if up_dir.exists():
                for f in sorted(up_dir.iterdir()):
                    if f.is_file() and not f.name.startswith('.'):
                        files.append({"name": f.name, "size": f.stat().st_size})
            self._json(200, {"files": files})

        elif p.startswith("/dashboard/") or p == "/dashboard":
            rel = p.removeprefix("/dashboard/") or "index.html"
            file_path = DASHBOARD_DIR / rel
            if file_path.exists() and file_path.is_file():
                ct = _mime(file_path.suffix)
                data = file_path.read_bytes()
                self.send_response(200)
                self.send_header("Content-Type", ct)
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)
            else:
                self.send_error(404)
        else:
            self.send_error(404)

    def do_POST(self):
        p = self.path.split("?")[0]
        length = int(self.headers.get("Content-Length", 0))
        body   = self.rfile.read(length).decode() if length else "{}"
        try:
            data = json.loads(body)
        except Exception:
            self._json(400, {"error": "bad json"}); return

        if p == "/upload-file":
            self._handle_file_upload(); return

        elif p == "/saved-runs":
            # Save a new run entry
            entry = data if isinstance(data, dict) else {}
            if not entry.get("id") or not entry.get("report"):
                self._json(400, {"error": "missing id or report"}); return
            runs = _load_saved_runs().get("runs", [])
            # Deduplicate by id
            runs = [r for r in runs if r.get("id") != entry["id"]]
            runs.insert(0, entry)
            if len(runs) > 50:
                runs = runs[:50]
            _save_saved_runs({"runs": runs})
            self._json(200, {"status": "ok", "count": len(runs)})

        elif p == "/run-analysis":
            with _state_lock:
                if _run_state["running"]:
                    self._json(200, {"status": "already_running"}); return
            t = threading.Thread(target=_run_worker, args=(data,), daemon=True)
            t.start()
            self._json(200, {"status": "started"})

        elif p == "/stop-analysis":
            import ctypes
            with _state_lock:
                is_running = _run_state["running"]
            if is_running:
                _stop_flag.set()
                # 1. Kill Ollama runner to unblock any in-flight LLM call
                _kill_ollama_now()
                # 2. Inject async exception into the worker thread
                tid = _worker_thread_id
                if tid:
                    try:
                        ctypes.pythonapi.PyThreadState_SetAsyncExc(
                            ctypes.c_ulong(tid),
                            ctypes.py_object(SystemExit),  # SystemExit is harder to swallow
                        )
                    except Exception:
                        pass
                # 3. Keep injecting every 500ms in background until thread dies
                wt = _worker_thread
                def _force_stop(tid=tid, wt=wt):
                    for _ in range(20):   # up to 10s
                        time.sleep(0.5)
                        if wt is None or not wt.is_alive():
                            break
                        if tid:
                            try:
                                ctypes.pythonapi.PyThreadState_SetAsyncExc(
                                    ctypes.c_ulong(tid),
                                    ctypes.py_object(SystemExit),
                                )
                            except Exception:
                                pass
                    # Final state cleanup if thread still hasn't ended
                    with _state_lock:
                        if _run_state["running"]:
                            _run_state["running"] = False
                            _run_state["logs"].append("[STOP] Force-stopped by user.")
                            for aid in _run_state["agent_states"]:
                                _run_state["agent_states"][aid] = "idle"
                threading.Thread(target=_force_stop, daemon=True).start()
                self._json(200, {"status": "stopping"})
            else:
                self._json(200, {"status": "not_running"})
        else:
            self.send_error(404)

    def _handle_file_upload(self):
        """Handle multipart file uploads. Saves to data/uploads/ and extracts text."""
        content_type = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in content_type:
            self._json(400, {"error": "Expected multipart/form-data"}); return

        up_dir = Path("data/uploads")
        up_dir.mkdir(parents=True, exist_ok=True)
        parsed_dir = Path("data/parsed")
        parsed_dir.mkdir(parents=True, exist_ok=True)

        try:
            form = cgi.FieldStorage(
                fp=self.rfile,
                headers=self.headers,
                environ={"REQUEST_METHOD": "POST", "CONTENT_TYPE": content_type},
            )
            uploaded = []
            for field_name in form.keys():
                item = form[field_name]
                if not hasattr(item, "filename") or not item.filename:
                    continue
                filename = Path(item.filename).name  # strip path
                save_path = up_dir / filename
                raw_bytes  = item.file.read()
                save_path.write_bytes(raw_bytes)

                # Try to extract text for agent context
                ext = save_path.suffix.lower()
                text_content = _extract_file_text(save_path, ext, raw_bytes)
                if text_content:
                    parsed_path = parsed_dir / (filename + ".txt")
                    parsed_path.write_text(text_content, encoding="utf-8")

                uploaded.append({
                    "name":     filename,
                    "size":     len(raw_bytes),
                    "parsed":   bool(text_content),
                    "ext":      ext,
                })

            self._json(200, {"status": "ok", "files": uploaded})
        except Exception as e:
            self._json(500, {"error": str(e)})

    def do_DELETE(self):
        p = self.path.split("?")[0]
        if p == "/saved-runs":
            # Clear all saved runs
            _save_saved_runs({"runs": []})
            self._json(200, {"status": "ok"})
        elif p.startswith("/saved-runs/"):
            # Delete single run by id
            run_id = p.removeprefix("/saved-runs/")
            try:
                run_id = int(run_id)
            except ValueError:
                self._json(400, {"error": "invalid id"}); return
            runs = [r for r in _load_saved_runs().get("runs", []) if r.get("id") != run_id]
            _save_saved_runs({"runs": runs})
            self._json(200, {"status": "ok", "count": len(runs)})
        else:
            self.send_error(404)

    def do_OPTIONS(self):
        self.send_response(200)
        self._cors()
        self.end_headers()

    def _json(self, code: int, payload: dict):
        body = json.dumps(payload).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self._cors()
        self.end_headers()
        self.wfile.write(body)

    def _cors(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, DELETE, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def log_message(self, *_):
        pass  # suppress request logging


def _mime(suffix: str) -> str:
    return {
        ".html": "text/html; charset=utf-8",
        ".css":  "text/css",
        ".js":   "application/javascript",
        ".json": "application/json",
        ".png":  "image/png",
        ".svg":  "image/svg+xml",
    }.get(suffix.lower(), "application/octet-stream")


def start(port: int = PORT):
    httpd = _ThreadingServer((HOST, port), Handler)
    t = threading.Thread(target=httpd.serve_forever, daemon=True, name="http-server")
    t.start()
    print(f"Hermes running at http://{HOST}:{port}/dashboard/index.html", flush=True)
    return httpd


if __name__ == "__main__":
    import webbrowser
    httpd = start()
    webbrowser.open(f"http://{HOST}:{PORT}/dashboard/index.html")
    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        print("\nShutting down.")
        httpd.shutdown()
